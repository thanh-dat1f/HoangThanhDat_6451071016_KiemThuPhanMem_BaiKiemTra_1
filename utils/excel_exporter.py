import os
import sys
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def export_test_cases_to_excel(results: list, output_filepath: str='reports/BaoCao_KiemThu_UTC.xlsx'):
    os.makedirs(os.path.dirname(output_filepath), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ket_Qua_Kiem_Thu_UTC'
    ws.views.sheetView[0].showGridLines = True
    font_title = Font(name='Times New Roman', size=15, bold=True, color='000000')
    font_header = Font(name='Times New Roman', size=13, bold=True, color='000000')
    font_body = Font(name='Times New Roman', size=13, bold=False, color='000000')
    font_sub = Font(name='Times New Roman', size=13, italic=True, color='000000')
    font_status_pass = Font(name='Times New Roman', size=13, bold=True, color='000000')
    font_status_fail = Font(name='Times New Roman', size=13, bold=True, color='000000')
    white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    border_black = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'), top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
    ws.merge_cells('A1:I1')
    ws['A1'] = 'BÁO CÁO KIỂM THỬ TỰ ĐỘNG (AUTOMATION TESTING REPORT)'
    ws['A1'].font = font_title
    ws['A1'].fill = white_fill
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40
    ws.merge_cells('A2:I2')
    ws['A2'] = 'Hệ thống: Văn phòng điện tử UTC (https://vanphongdientu.utc.edu.vn/) | Framework: Selenium WebDriver + Python'
    ws['A2'].font = font_sub
    ws['A2'].fill = white_fill
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 25
    total_tests = len(results)
    pass_count = sum((1 for r in results if r.get('status') == 'PASS'))
    fail_count = total_tests - pass_count
    pass_rate = round(pass_count / total_tests * 100, 1) if total_tests > 0 else 0
    stats = [('Tổng số Test Cases:', total_tests, 'A4', 'B4'), ('Số ca kiểm thử PASS:', pass_count, 'C4', 'D4'), ('Số ca kiểm thử FAIL:', fail_count, 'E4', 'F4'), ('Tỷ lệ đạt (Pass Rate):', f'{pass_rate}%', 'G4', 'H4')]
    for label, val, c1, c2 in stats:
        ws[c1] = label
        ws[c1].font = font_header
        ws[c1].fill = white_fill
        ws[c1].alignment = Alignment(horizontal='right', vertical='center')
        ws[c1].border = border_black
        ws[c2] = val
        ws[c2].font = font_header
        ws[c2].fill = white_fill
        ws[c2].alignment = Alignment(horizontal='center', vertical='center')
        ws[c2].border = border_black
    ws.row_dimensions[4].height = 30
    headers = ['Mã TC', 'Tên Kịch Bản Kiểm Thử', 'Kỹ Thuật Kiểm Thử', 'Dữ Liệu Thử Nghiệm (Test Data)', 'Các Bước Thực Hiện', 'Kết Quả Mong Đợi', 'Kết Quả Thực Tế', 'Trạng Thái', 'Ảnh Minh Chứng']
    header_row = 6
    ws.row_dimensions[header_row].height = 35
    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.fill = white_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border_black
    current_row = 7
    for r in results:
        status = r.get('status', 'FAIL')
        screenshot = r.get('screenshot', '')
        screenshot_display = os.path.basename(screenshot) if screenshot else 'Không có'
        row_data = [r.get('id', ''), r.get('name', ''), r.get('technique', 'Kiểm thử chức năng'), r.get('test_data', '-'), r.get('steps', '-'), r.get('expected', '-'), r.get('actual', '-'), status, screenshot_display]
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_body
            cell.fill = white_fill
            cell.border = border_black
            if col_idx in [1, 8, 9]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            if col_idx == 8:
                cell.font = font_status_pass if status == 'PASS' else font_status_fail
        ws.row_dimensions[current_row].height = 55
        current_row += 1
    col_widths = {1: 16, 2: 34, 3: 26, 4: 28, 5: 38, 6: 34, 7: 34, 8: 16, 9: 24}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    try:
        wb.save(output_filepath)
        final_path = output_filepath
    except PermissionError:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        dir_name = os.path.dirname(output_filepath)
        base_name = os.path.splitext(os.path.basename(output_filepath))[0]
        final_path = os.path.join(dir_name, f'{base_name}_{timestamp}.xlsx')
        wb.save(final_path)
        print(f"[CẢNH BÁO]: File '{output_filepath}' đang được mở trong Microsoft Excel.")
        print(f'[ĐÃ LƯU BẢN THAY THẾ]: {os.path.abspath(final_path)}')
    print(f'[XUẤT EXCEL THÀNH CÔNG]: {os.path.abspath(final_path)}')
    return os.path.abspath(final_path)
